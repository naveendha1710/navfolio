# backend/main.py
from fastapi import FastAPI, Form, Body, Request
from fastapi.middleware.cors import CORSMiddleware
from datetime import datetime
from openpyxl import load_workbook, Workbook
from email.message import EmailMessage
from pydantic import BaseModel, EmailStr
import smtplib
import os

app = FastAPI(title="Portfolio Lead Collector & Resume Sender")

# Configure CORS for the React frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://localhost:5173",
        "http://127.0.0.1:3000",
        "http://127.0.0.1:5173",
        "http://192.168.1.7:5173"
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

OUTLOOK_EMAIL = "nav.cs@outlook.com"
OUTLOOK_APP_PASSWORD = os.getenv("OUTLOOK_APP_PASSWORD", "YOUR_16_CHAR_MICROSOFT_APP_PASSWORD")

EXCEL_FILE = os.path.join(os.path.dirname(__file__), "portfolio_leads.xlsx")
RESUME_PATH = os.path.join(os.path.dirname(__file__), "Resume.pdf")

class ContactPayload(BaseModel):
    name: str = ""
    email: str
    message: str = ""
    _gotcha: str = ""

def append_to_excel(user_name: str, user_email: str, user_message: str):
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Leads"
        ws.append(["Timestamp", "Name", "Email", "Message"])
        wb.save(EXCEL_FILE)

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), user_name, user_email, user_message])
    wb.save(EXCEL_FILE)

def send_resume_email(recipient_email: str, recipient_name: str = ""):
    """Sends resume email using Outlook SMTP."""
    msg = EmailMessage()
    msg['Subject'] = "Thanks for connecting! | My Portfolio & Resume"
    msg['From'] = f"Nav <{OUTLOOK_EMAIL}>"
    msg['To'] = recipient_email
    
    greeting = f"Hi {recipient_name}," if recipient_name else "Hi there,"
    body = f"""{greeting}

Thanks for checking out my portfolio! As requested, I have attached my resume to this email. 

Let me know if you have any questions or would like to schedule a quick chat regarding any development opportunities.

Best regards,
Nav"""
    msg.set_content(body)

    if os.path.exists(RESUME_PATH):
        with open(RESUME_PATH, "rb") as f:
            msg.add_attachment(
                f.read(),
                maintype='application',
                subtype='pdf',
                filename="Nav_Resume.pdf"
            )

    with smtplib.SMTP("smtp-mail.outlook.com", 587) as server:
        server.starttls()
        server.login(OUTLOOK_EMAIL, OUTLOOK_APP_PASSWORD)
        server.send_message(msg)

@app.get("/")
def read_root():
    return {"status": "online", "message": "Portfolio Lead API is running."}

@app.post("/api/contact")
async def handle_contact(request: Request):
    try:
        # Support both JSON and Form payloads
        content_type = request.headers.get("content-type", "")
        if "application/json" in content_type:
            data = await request.json()
            name = data.get("name", "")
            email = data.get("email", "")
            message = data.get("message", "")
        else:
            form = await request.form()
            name = form.get("name", "")
            email = form.get("email", "")
            message = form.get("message", "")

        if not email:
            return {"status": "error", "message": "Email is required."}

        append_to_excel(name, email, message)

        try:
            send_resume_email(email, name)
            email_sent = True
        except Exception as mail_err:
            print(f"[Warning] Excel logged, but email sending failed: {mail_err}")
            email_sent = False

        return {
            "status": "success",
            "message": "Thank you! Your message has been sent successfully." if email_sent else "Email logged successfully.",
            "email_sent": email_sent
        }
    except Exception as e:
        return {"status": "error", "message": str(e)}
